"""Text extraction: PDF text layer first (PyMuPDF), RapidOCR (ONNX) for scans and images.

extract_text(path) -> {
  text:        all pages joined with "\f" (form feed) - page i is text.split("\f")[i]
  pages:       page count (int)            pageCount: same (alias used by Document.ocr)
  pageTexts:   [str]                       pageSizes: [[width, height]] (PDF points or image pixels)
  method:      "pdf-text" | "rapidocr" | "pdf-text+rapidocr" | "text" (csv/txt/xlsx) | "none"
  confidence:  0..1
  blocks:      [{text, bbox:[x0,y0,x1,y1], page (1-based), conf}]   one block per visual text line
}
Lines are re-assembled into reading-order rows (cells on the same baseline joined by two spaces) so that
table rows and "Label: value" pairs survive as single lines for the extractor.
"""
from __future__ import annotations

import csv
import gc
import io
import threading
from pathlib import Path

import numpy as np

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
TEXT_EXT = {".csv", ".tsv", ".txt"}
XLSX_EXT = {".xlsx", ".xlsm"}
MIN_TEXT_CHARS = 40       # below this a PDF page is treated as a scan and OCR'd
OCR_ZOOM = 2.0            # rasterisation zoom for scanned PDF pages (144 dpi)
MAX_OCR_SIDE = 2200

_engine = None
_lock = threading.Lock()
_inference_lock = threading.Lock()


def get_ocr_engine():
    """Lazily construct RapidOCR; optional caching is for larger instances."""
    global _engine
    with _lock:
        if _engine is None:
            from rapidocr_onnxruntime import RapidOCR
            from app.config import get_settings

            # Hosts can expose many CPU cores despite granting a small CPU/RAM
            # quota. Bound each ONNX session instead of auto-sizing thread pools.
            threads = max(1, int(get_settings().ocr_threads))
            _engine = RapidOCR(intra_op_num_threads=threads, inter_op_num_threads=threads,
                               rec_batch_num=1, cls_batch_num=1,
                               det_limit_type="max", det_limit_side_len=960)
    return _engine


def ocr_image_array(img_bgr: np.ndarray, page: int = 1, scale: float = 1.0) -> list[dict]:
    """Run RapidOCR on a BGR array. bboxes are divided by `scale` (to map back to page units)."""
    import cv2

    h, w = img_bgr.shape[:2]
    factor = 1.0
    if max(h, w) > MAX_OCR_SIDE:
        factor = MAX_OCR_SIDE / max(h, w)
        img_bgr = cv2.resize(img_bgr, (int(w * factor), int(h * factor)), interpolation=cv2.INTER_AREA)
    # Create/use/release under one lock: model construction itself is expensive.
    # On small hosts, don't retain model memory through ELA/QR analysis afterward.
    with _inference_lock:
        from app.config import get_settings

        engine = get_ocr_engine()
        try:
            result, _ = engine(img_bgr)
        finally:
            if not get_settings().ocr_cache:
                global _engine
                with _lock:
                    _engine = None
                del engine
                gc.collect()
    blocks = []
    for box, text, score in result or []:
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        k = scale * factor
        blocks.append({"text": str(text), "bbox": [round(min(xs) / k, 1), round(min(ys) / k, 1),
                                                   round(max(xs) / k, 1), round(max(ys) / k, 1)],
                       "page": page, "conf": round(float(score), 3)})
    return blocks


def rows_from_blocks(blocks: list[dict]) -> list[dict]:
    """Group line blocks of ONE page into visual rows (vertical overlap), left-to-right within a row."""
    items = sorted(blocks, key=lambda b: ((b["bbox"][1] + b["bbox"][3]) / 2, b["bbox"][0]))
    rows: list[list[dict]] = []
    for b in items:
        y0, y1 = b["bbox"][1], b["bbox"][3]
        cy = (y0 + y1) / 2
        placed = False
        if rows:
            last = rows[-1]
            ly0 = min(x["bbox"][1] for x in last)
            ly1 = max(x["bbox"][3] for x in last)
            lh = max(1.0, min(ly1 - ly0, y1 - y0))
            if ly0 - 0.3 * lh <= cy <= ly1 + 0.3 * lh and abs(cy - (ly0 + ly1) / 2) < 0.6 * max(lh, 1.0):
                last.append(b)
                placed = True
        if not placed:
            rows.append([b])
    out = []
    for r in rows:
        r.sort(key=lambda b: b["bbox"][0])
        out.append({
            "text": "  ".join(x["text"].strip() for x in r if x["text"].strip()),
            "bbox": [min(x["bbox"][0] for x in r), min(x["bbox"][1] for x in r),
                     max(x["bbox"][2] for x in r), max(x["bbox"][3] for x in r)],
            "page": r[0]["page"],
            "conf": round(sum(x.get("conf", 1.0) for x in r) / len(r), 3),
        })
    return out


def _pdf_page_lines(page, page_no: int) -> list[dict]:
    d = page.get_text("dict")
    lines = []
    for blk in d.get("blocks", []):
        if blk.get("type") != 0:
            continue
        for ln in blk.get("lines", []):
            text = "".join(sp.get("text", "") for sp in ln.get("spans", []))
            if text.strip():
                lines.append({"text": text, "bbox": [round(v, 1) for v in ln["bbox"]], "page": page_no, "conf": 1.0})
    return lines


def _extract_pdf(path: Path) -> dict:
    import cv2
    import pymupdf

    page_texts, blocks, sizes, methods, confs = [], [], [], set(), []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc):
            pno = i + 1
            sizes.append([round(page.rect.width, 1), round(page.rect.height, 1)])
            lines = _pdf_page_lines(page, pno)
            if sum(len(l["text"].strip()) for l in lines) >= MIN_TEXT_CHARS:
                methods.add("pdf-text")
                confs.append(0.99)
            else:
                pix = page.get_pixmap(matrix=pymupdf.Matrix(OCR_ZOOM, OCR_ZOOM))
                arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
                arr = cv2.cvtColor(arr, cv2.COLOR_RGBA2BGR if pix.n == 4 else cv2.COLOR_RGB2BGR) if pix.n >= 3 \
                    else cv2.cvtColor(arr, cv2.COLOR_GRAY2BGR)
                lines = ocr_image_array(arr, page=pno, scale=OCR_ZOOM)
                methods.add("rapidocr")
                if lines:
                    confs.append(sum(l["conf"] for l in lines) / len(lines))
            rows = rows_from_blocks(lines)
            blocks.extend(rows)
            page_texts.append("\n".join(r["text"] for r in rows))
    if methods == {"pdf-text"}:
        method = "pdf-text"
    elif methods == {"rapidocr"}:
        method = "rapidocr"
    elif methods:
        method = "pdf-text+rapidocr"
    else:
        method = "none"
    return _result(page_texts, blocks, sizes, method, sum(confs) / len(confs) if confs else 0.0)


def _extract_image(path: Path) -> dict:
    import cv2

    img = cv2.imdecode(np.fromfile(str(path), dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        return _result([""], [], [[0, 0]], "none", 0.0)
    lines = ocr_image_array(img, page=1)
    rows = rows_from_blocks(lines)
    conf = sum(l["conf"] for l in lines) / len(lines) if lines else 0.0
    return _result(["\n".join(r["text"] for r in rows)], rows, [[img.shape[1], img.shape[0]]], "rapidocr", conf)


def _extract_tabular(path: Path) -> dict:
    ext = path.suffix.lower()
    if ext in XLSX_EXT:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in row):
                    w.writerow(["" if c is None else str(c) for c in row])
        text = buf.getvalue()
    else:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    lines = text.splitlines()
    blocks = [{"text": l, "bbox": [0, i, 0, i + 1], "page": 1, "conf": 1.0} for i, l in enumerate(lines) if l.strip()]
    return _result([text.replace("\f", "\n")], blocks, [[0, 0]], "text", 1.0)


def _result(page_texts, blocks, sizes, method, conf) -> dict:
    return {
        "text": "\f".join(page_texts),
        "pages": len(page_texts),
        "pageCount": len(page_texts),
        "pageTexts": page_texts,
        "pageSizes": sizes,
        "method": method,
        "confidence": round(float(conf), 3),
        "blocks": blocks,
    }


def extract_text(path: str | Path) -> dict:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext in IMAGE_EXT:
        return _extract_image(path)
    if ext in TEXT_EXT or ext in XLSX_EXT:
        return _extract_tabular(path)
    return _result([""], [], [[0, 0]], "none", 0.0)
